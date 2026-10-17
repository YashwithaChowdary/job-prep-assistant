"""Day 27: turn a resume into structured data.

Pipeline: file -> text -> redact personal data -> sections -> skills (with
evidence), education, projects, years of experience.
"""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from .config import DEFAULT, Config
from .matcher import find_mentions
from .skills_kb import KB

# --------------------------------------------------------------- file -> text


def read_text(path: str | Path) -> str:
    """Read .txt/.md, .docx or .pdf. Raises ValueError with a clear message otherwise."""
    p = Path(path)
    ext = p.suffix.lower()
    if ext in {".txt", ".md", ""}:
        return p.read_text(encoding="utf-8", errors="replace")
    if ext == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".pdf":
        from pypdf import PdfReader
        text = "\n".join((page.extract_text() or "") for page in PdfReader(str(p)).pages)
        if len(text.strip()) < 40:
            raise ValueError(f"{p.name}: no text found. It may be a scanned image; OCR is not supported.")
        return text
    raise ValueError(f"Unsupported file type: {ext}. Use .txt, .docx or .pdf.")


_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_URL = re.compile(r"(?:https?://|www\.)\S+|\b(?:linkedin|github)\.com/\S+", re.I)
_PHONE = re.compile(r"(?<![\w])\+?(?:\d[\s().-]?){10,14}(?![\w])")


def redact_pii(text: str) -> str:
    """Remove emails, phone numbers and profile links before anything else sees the text."""
    return _PHONE.sub("[PHONE]", _URL.sub("[URL]", _EMAIL.sub("[EMAIL]", text)))


# ------------------------------------------------------------------ sections
_HEAD = {
    "summary": {"summary", "profile", "objective", "about me", "professional summary", "career objective"},
    "skills": {"skills", "technical skills", "core competencies", "key skills", "tools", "technologies",
               "skills and tools", "tech stack"},
    "experience": {"experience", "work experience", "professional experience", "employment",
                   "work history", "internships"},
    "projects": {"projects", "academic projects", "personal projects", "key projects", "selected projects"},
    "education": {"education", "academics", "academic background", "qualifications"},
    "certifications": {"certifications", "certificates", "courses", "training"},
    "other": {"achievements", "awards", "publications", "interests", "languages", "references"},
}
_HEAD_LOOKUP = {h: sec for sec, hs in _HEAD.items() for h in hs}
DEMONSTRATED = {"experience", "projects", "summary"}   # places that show skills in use


def _heading(line: str) -> str | None:
    s = line.strip().rstrip(":").strip()
    if not s or len(s) > 40:
        return None
    key = re.sub(r"[^a-z ]", "", s.lower()).strip()
    return _HEAD_LOOKUP.get(key)


def split_sections(text: str) -> list[tuple[str, str]]:
    """Return [(section, line)] in order. Text before the first heading is 'summary'."""
    out, current = [], "summary"
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        sec = _heading(line)
        if sec:
            current = sec
            continue
        m = re.match(r"(?i)^(technical skills|skills)\s*:\s*(.*)$", line)
        if m:                                           # inline "Skills: a, b, c"
            out.append(("skills", m.group(2)))
            continue
        out.append((current, line))
    return out


# -------------------------------------------------------------------- skills

def extract_skills(sections: list[tuple[str, str]], cfg: Config) -> tuple[dict, list[str]]:
    skills: dict[str, dict] = {}
    negated: set[str] = set()
    for section, line in sections:
        for m in find_mentions(line, cfg):
            if m.level == "negated":
                negated.add(m.skill)
                continue
            rec = skills.setdefault(m.skill, {"level": "partial", "sections": [], "evidence": []})
            if m.level == "full":
                rec["level"] = "full"
            if section not in rec["sections"]:
                rec["sections"].append(section)
            if len(rec["evidence"]) < 3:
                rec["evidence"].append({"section": section, "text": _clip(line)})
    for rec in skills.values():
        rec["strength"] = "demonstrated" if set(rec["sections"]) & DEMONSTRATED else "listed"
    negated -= set(skills)
    return skills, sorted(negated)


def add_implied(skills: dict[str, dict], cfg: Config) -> None:
    if not cfg.use_implied:
        return
    for name, rec in list(skills.items()):
        if rec["level"] != "full":
            continue
        for target in KB[name].implies:
            if target not in skills:
                skills[target] = {"level": "full", "strength": "implied", "sections": [],
                                  "evidence": [{"section": "implied", "text": f"Implied by {name}"}],
                                  "implied_by": [name]}
            elif skills[target].get("strength") == "implied" and name not in skills[target]["implied_by"]:
                skills[target]["implied_by"].append(name)
                skills[target]["evidence"][0]["text"] = "Implied by " + ", ".join(skills[target]["implied_by"])


def _clip(s: str, n: int = 110) -> str:
    s = re.sub(r"\s+", " ", s).strip()
    return s if len(s) <= n else s[: n - 1] + "…"


def unrecognised_skills(sections: list[tuple[str, str]], cfg: Config) -> list[str]:
    """Items listed in a Skills section that the knowledge base does not know."""
    out: list[str] = []
    for section, line in sections:
        if section != "skills":
            continue
        for item in re.split(r"[,;|•/]", line):
            item = re.sub(r"\(.*?\)", "", item).strip(" -–*·")
            if 2 <= len(item) <= 30 and len(item.split()) <= 4 and not find_mentions(item, cfg, with_cues=False):
                out.append(item)
    return out


# ----------------------------------------------------------------- education
_DEGREE = re.compile(
    r"\b(B\.?\s?Tech|B\.?E\.?|B\.?Sc|B\.?Com|BCA|BBA|BFA|B\.?A\.?|Bachelor\w*|M\.?\s?Tech|M\.?E\.?|M\.?Sc|MCA|MBA|"
    r"M\.?A\.?|Master\w*|Ph\.?D|Diploma)\b")


def extract_education(sections) -> list[dict]:
    out = []
    for section, line in sections:
        if section == "education" and _DEGREE.search(line):
            years = re.findall(r"(?:19|20)\d{2}", line)
            out.append({"text": _clip(line, 140), "degree": _DEGREE.search(line).group(1),
                        "year": years[-1] if years else None})
    return out


# ------------------------------------------------------------------ projects
_BULLET = re.compile(r"^\s*[-•*–·]\s*")


def extract_projects(sections, cfg: Config) -> list[dict]:
    blocks: list[dict] = []
    for section, line in sections:
        if section != "projects":
            continue
        if _BULLET.match(line) and blocks and blocks[-1]["has_title"]:
            blocks[-1]["lines"].append(_BULLET.sub("", line))
        else:
            clean = _BULLET.sub("", line)
            title = re.split(r"\s[-–—|]\s|:\s", clean, maxsplit=1)[0]
            blocks.append({"title": _clip(title, 60), "lines": [clean], "has_title": not _BULLET.match(line) or True})
    out = []
    for b in blocks:
        text = " ".join(b["lines"])
        names = []
        for m in find_mentions(text, cfg):
            if m.level != "negated" and m.skill not in names:
                names.append(m.skill)
        out.append({"title": b["title"], "description": _clip(text, 200), "skills": names})
    return out


# ---------------------------------------------------------------- experience
_MON = "jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec"
_MONTH_NUM = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
_RANGE = re.compile(
    rf"(?:(?P<m1>{_MON})[a-z]*\.?\s+)?(?P<y1>(?:19|20)\d{{2}})\s*(?:[-–—]|\bto\b)\s*"
    rf"(?:(?:(?P<m2>{_MON})[a-z]*\.?\s+)?(?P<y2>(?:19|20)\d{{2}})|(?P<now>present|current|now|till date|ongoing))", re.I)


def _month(m: str | None, default: int) -> int:
    return _MONTH_NUM.get(m.lower()[:3], default) if m else default


def years_of_experience(sections, today: date) -> tuple[float | None, list[str]]:
    spans, roles = [], []
    for section, line in sections:
        if section != "experience":
            continue
        for m in _RANGE.finditer(line):
            start = int(m["y1"]) * 12 + _month(m["m1"], 1)
            end = today.year * 12 + today.month if m["now"] else int(m["y2"]) * 12 + _month(m["m2"], 12)
            if end >= start:
                spans.append((start, end))
                roles.append(_clip(line, 100))
    if not spans:
        return None, []
    spans.sort()
    merged = [list(spans[0])]
    for s, e in spans[1:]:
        if s <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], e)
        else:
            merged.append([s, e])
    months = sum(e - s + 1 for s, e in merged)
    return round(months / 12, 1), roles


# ---------------------------------------------------------------------- main

def extract_resume(text: str, cfg: Config = DEFAULT, today: date | None = None) -> dict:
    today = today or date.today()
    text = redact_pii(text)
    sections = split_sections(text)
    skills, negated = extract_skills(sections, cfg)
    add_implied(skills, cfg)
    years, roles = years_of_experience(sections, today)
    present = {s for s, _ in sections}
    warnings = []
    if "skills" not in present:
        warnings.append("No Skills section found; skills were read from the whole text.")
    if "experience" not in present:
        warnings.append("No Experience section found; years of experience unknown.")
    elif years is None:
        warnings.append("Experience section has no readable dates; years of experience unknown.")
    return {
        "skills": skills,
        "negated_skills": negated,
        "unrecognised_skills": unrecognised_skills(sections, cfg),
        "education": extract_education(sections),
        "projects": extract_projects(sections, cfg),
        "years_experience": years,
        "roles": roles,
        "sections_found": sorted(present),
        "warnings": warnings,
    }
